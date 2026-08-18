#!/usr/bin/env python3
"""Extract insighter MVP blocks JSON from insighter_outer xlsx (stdout)."""
import json
import re
import sys

import openpyxl

PURE_COLS = [
    "individualist",
    "innovator",
    "influencer",
    "fashionista",
    "conservative",
    "sensitive",
    "reserved",
    "realist",
]

LEGACY = {
    "fashionista": "aesthete",
    "sensitive": "empath",
    "reserved": "analyst",
    "realist": "pragmatist",
}

# TZ v4.0 §4.2 — 10 MVP block_key for mode insighter
ROW_DEFS = [
    {
        "block_key": "burnout_reason",
        "group": "Выгорание и стресс",
        "index": 0,
        "sort_order": 10,
        "title": "Выгорание и стресс",
    },
    {
        "block_key": "how_to_deal_with_burnout",
        "group": "Выгорание и стресс",
        "index": 1,
        "sort_order": 20,
        "title": "Как справляться с выгоранием",
    },
    {
        "block_key": "what_drives_and_inspires",
        "group": "Дисциплина и мотивация",
        "index": 0,
        "sort_order": 30,
        "title": "Что мотивирует",
    },
    {
        "block_key": "what_triggers_and_demotivates",
        "group": "Дисциплина и мотивация",
        "index": 1,
        "sort_order": 40,
        "title": "Что демотивирует",
    },
    {
        "block_key": "how_to_praise",
        "group": "Фидбек",
        "index": 0,
        "sort_order": 50,
        "title": "Как хвалить",
    },
    {
        "block_key": "how_to_give_negative_feedback",
        "group": "Фидбек",
        "index": 1,
        "sort_order": 60,
        "title": "Как давать негативный фидбек",
    },
    {
        "block_key": "how_to_onboard",
        "group": "Коммуникация",
        "index": 0,
        "sort_order": 70,
        "title": "Как онбордить",
    },
    {
        "block_key": "team_player",
        "group": "Отношения в команде",
        "index": 0,
        "sort_order": 80,
        "title": "Роль в команде",
    },
    {
        "block_key": "how_conflicting",
        "group": "Отношения в команде",
        "index": 1,
        "sort_order": 90,
        "title": "Поведение в конфликте",
    },
    {
        "block_key": "how_to_resolve_a_conflict",
        "group": "Отношения в команде",
        "index": 2,
        "sort_order": 100,
        "title": "Как разрешать конфликт",
    },
]


def normalize_group(value: str) -> str:
    text = str(value).strip().split("\n")[0]
    text = re.sub(r"\s*\([a-zA-Z]+\)\s*$", "", text).strip()
    return text


def main() -> None:
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = "Результат (новый)" if "Результат (новый)" in wb.sheetnames else wb.sheetnames[0]
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = rows[0]
    col_by_code: dict[str, int] = {}
    for idx, cell in enumerate(header):
        if not cell:
            continue
        text = str(cell).replace("\n", " ")
        match = re.search(r"\(([a-z_]+)\)\s*$", text)
        if not match:
            continue
        code = match.group(1)
        if code in PURE_COLS:
            col_by_code[code] = idx

    grouped: dict[str, list[tuple]] = {}
    current_group: str | None = None
    for row in rows[1:]:
        if row[0]:
            current_group = normalize_group(str(row[0]))
        if not current_group:
            continue
        grouped.setdefault(current_group, []).append(row)

    blocks = []
    for defn in ROW_DEFS:
        group_rows = grouped.get(defn["group"], [])
        if defn["index"] >= len(group_rows):
            raise SystemExit(
                f"Missing row {defn['index']} for group {defn['group']} ({defn['block_key']})"
            )
        row = group_rows[defn["index"]]
        for legacy_code, col_idx in col_by_code.items():
            psychotype = LEGACY.get(legacy_code, legacy_code)
            text = row[col_idx] if col_idx < len(row) else None
            if text is None:
                continue
            content = str(text).strip()
            if not content or content.lower() == "none":
                continue
            blocks.append(
                {
                    "psychotype_code": psychotype,
                    "mode": "insighter",
                    "block_key": defn["block_key"],
                    "block_template": "text_block",
                    "title": defn["title"],
                    "content_text": content,
                    "sort_order": defn["sort_order"],
                    "is_visible_default": True,
                    "version": 1,
                    "is_active": True,
                }
            )

    print(json.dumps({"blocks": blocks}, ensure_ascii=False))


if __name__ == "__main__":
    main()
