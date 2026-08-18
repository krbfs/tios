#!/usr/bin/env python3
"""Extract insighter_you blocks JSON from source xlsx (stdout)."""
import json
import re
import sys

import openpyxl

PURE_COLS = [
    "individualist",
    "innovator",
    "influencer",
    "fashionista",
    "conservative",
    "sensitive",
    "reserved",
    "realist",
]

LEGACY = {
    "fashionista": "aesthete",
    "sensitive": "empath",
    "reserved": "analyst",
    "realist": "pragmatist",
}

ROW_DEFS = [
    {"block_key": "teamPlayer", "sort_order": 10},
    {"block_key": "productivityDependsOn", "sort_order": 20},
    {"block_key": "typesOfWorkflows", "sort_order": 30},
    {"block_key": "whatTriggersAndDemotivates", "sort_order": 40},
    {"block_key": "howConflicting", "sort_order": 50},
    {"block_key": "independenceAndSelfOrganized", "sort_order": 60},
    {"block_key": "thinkingStyleAndInfoDelivery", "sort_order": 70},
    {"block_key": "strengthsAtWork", "sort_order": 80},
    {"block_key": "burnoutReason", "sort_order": 90},
    {"block_key": "optimalTrainingSystem", "sort_order": 100},
]


def main() -> None:
    path = sys.argv[1]
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header = rows[0]
    col_by_code: dict[str, int] = {}
    for idx, cell in enumerate(header):
        if not cell:
            continue
        text = str(cell)
        match = re.search(r"\(([a-z_]+)\)\s*$", text, re.M)
        if not match:
            continue
        code = match.group(1)
        if code in PURE_COLS:
            col_by_code[code] = idx

    blocks = []
    data_rows = [row for row in rows[1:] if row and row[0]]

    for row_index, row in enumerate(data_rows[: len(ROW_DEFS)]):
        defn = ROW_DEFS[row_index]
        title = str(row[0]).strip()
        for legacy_code, col_idx in col_by_code.items():
            psychotype = LEGACY.get(legacy_code, legacy_code)
            text = row[col_idx] if col_idx < len(row) else None
            if text is None:
                continue
            content = str(text).strip()
            if not content or content.lower() == "none":
                continue
            blocks.append(
                {
                    "psychotype_code": psychotype,
                    "mode": "insighter_you",
                    "block_key": defn["block_key"],
                    "block_template": "text_block",
                    "title": title,
                    "content_text": content,
                    "sort_order": defn["sort_order"],
                    "is_visible_default": True,
                    "version": 1,
                    "is_active": True,
                }
            )

    print(json.dumps({"blocks": blocks}, ensure_ascii=False))


if __name__ == "__main__":
    main()
